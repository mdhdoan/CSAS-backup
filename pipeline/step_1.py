import json
import sys

from openpyxl import load_workbook

from langchain.output_parsers.list import NumberedListOutputParser
from langchain.prompts.prompt import PromptTemplate
from langchain.schema import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter

from langchain_community.llms import Ollama

from sentence_transformers import SentenceTransformer

import torch


QUESTIONS = [
    'What are the problems for the fish population?',
    'What are the risks for the fish population?',
]


def create_qa_prompt(format_instructions):
    QA_TEMPLATE = """Provide a short and concise answer for the following question using ONLY information found in the articles delimited by triple backquotes (```).
    Return answer with highest confidence score. Do not explain.
    QUESTION:{question}

    ARTICLE:```{article}```

    ANSWERS:{format_instructions}:"""
    return PromptTemplate(
        input_variables=["question", "article"], 
        partial_variables={"format_instructions": format_instructions}, 
        template=QA_TEMPLATE)


def answer_questions(llm_chain, text, questions):
    qa_dict = dict()    
    for question in questions:
        output = llm_chain.invoke({'article': text, 'question': question})
        qa_dict[question] = output[0]
    return qa_dict


def load_data_file(file_name, sheet_name):
    wb = load_workbook(file_name)
    ws = wb[sheet_name]
    headers, inputs = [], []
    for i, row in enumerate(ws): 
        if i == 0:
            headers = [cell.value for cell in row]
            continue
        inputs.append({headers[j]: cell.value for j, cell in enumerate(row)})
    return inputs


def load_jsonl(file_name, slice=None, single=False):
    with open(file_name, 'rt') as in_file:
        documents = dict()
        lines = in_file.readlines()[0:int(slice)] if slice else in_file.readlines()
        for d in [json.loads(line.strip()) for line in lines]:
            documents[d['metadata']['link']] = d
        print(f"[{file_name}] Read {len(documents)} documents.")
        return documents


if __name__ == '__main__':
    input_data_file = sys.argv[1]
    input_sheet_name = sys.argv[2]
    jsonl_data_file = sys.argv[3]
    llm_model_name = sys.argv[4]
    
    text_splitter = RecursiveCharacterTextSplitter(chunk_size=200, chunk_overlap=0)
    model = SentenceTransformer("all-mpnet-base-v2")

    output_parser = NumberedListOutputParser()
    format_instructions = output_parser.get_format_instructions()
    llm_model = Ollama(model=llm_model_name, temperature=0.0)
    qa_prompt = create_qa_prompt(format_instructions)
    qa_llm_chain = qa_prompt | llm_model | output_parser
    
    inputs = load_data_file(input_data_file, input_sheet_name)
    documents = load_jsonl(jsonl_data_file)
    
    for input in inputs:
        document = documents[input['Link']]
        print(f"\n{input['Link']}\n\t[SUM] {document['metadata']['summary']}")
        
        print(f"\t[QA SUM]")
        qa_dict = answer_questions(qa_llm_chain, document['metadata']['summary'], QUESTIONS)
        for question, answer in qa_dict.items():
            print(f"\t--- {question} --- {answer}")
        q1_embeddings = model.encode([qa_dict[QUESTIONS[0]]], convert_to_tensor=True)
        q2_embeddings = model.encode([qa_dict[QUESTIONS[1]]], convert_to_tensor=True)
        
        problem = input['Problem']
        p_embeddings = model.encode([problem], convert_to_tensor=True)
        similarity_scores = model.similarity(p_embeddings, q1_embeddings)[0]
        scores, indices = torch.topk(similarity_scores, k=1)
        print(f"\t[QA P SIM] -- {qa_dict[QUESTIONS[0]]}")
        for score, idx in zip(scores, indices):
            print(f"\t--- {score:.4f} --- {problem}")

        risk = input['Risk']
        r_embeddings = model.encode([risk], convert_to_tensor=True)
        similarity_scores = model.similarity(r_embeddings, q2_embeddings)[0]
        scores, indices = torch.topk(similarity_scores, k=1)
        print(f"\t[QA R SIM] -- {qa_dict[QUESTIONS[1]]}")
        for score, idx in zip(scores, indices):
            print(f"\t--- {score:.4f} --- {risk}")

        docs = text_splitter.split_documents([Document(page_content=document['page_content'])])
        t_embeddings = model.encode([d.page_content for d in docs], convert_to_tensor=True)
        
        similarity_scores = model.similarity(p_embeddings, t_embeddings)[0]
        scores, indices = torch.topk(similarity_scores, k=3)
        print(f"\t[PROB] -- {problem}")
        for score, idx in zip(scores, indices):
            print(f"\t--- {score:.4f} --- {docs[idx].page_content}")

        similarity_scores = model.similarity(r_embeddings, t_embeddings)[0]
        scores, indices = torch.topk(similarity_scores, k=3)
        print(f"\t[RISK] -- {risk}")
        for score, idx in zip(scores, indices):
            print(f"\t--- {score:.4f} --- {docs[idx].page_content}")

